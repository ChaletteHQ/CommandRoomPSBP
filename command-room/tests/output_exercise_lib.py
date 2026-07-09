#!/usr/bin/env python3
"""SPEC A8 — shared library for the output-skill regression exercises.

NOT a test (no run_/test_ prefix → not discovered by run_all.py). Provides the
fixture copy, .docx text extraction + normalization, golden-snapshot compare (with
CR_UPDATE_GOLDENS=1), a stdlib event validator that reads the live schema enum, and
the placeholder check. The per-skill exercises import these.
"""
from __future__ import annotations

import difflib
import json
import os
import re
import shutil
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "tests" / "fixtures" / "workspace_mini"
GOLDEN_DIR = ROOT / "tests" / "golden"
SCHEMA_PATH = ROOT / "shared" / "data-schemas" / "events.schema.json"

PLACEHOLDERS = ("TBD", "TODO", "Lorem", "[name]", "[topic]", "[Counterpart]",
                "<topic>", "XXX", "FIXME", "[date]")

_DATE_ISO = re.compile(r"\b\d{4}-\d{2}-\d{2}\b")
_DATE_LONG = re.compile(
    r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)"
    r"\s+\d{1,2},\s+\d{4}\b"
)


# ---- exercise scaffolding (mirrors runtime_exercise_research.py) ----

def make_recorder():
    PASS: list[str] = []
    FAIL: list[str] = []

    def ok(name: str, detail: str = "") -> None:
        PASS.append(name)
        print(f"  PASS  {name}" + (f"  ({detail})" if detail else ""))

    def fail(name: str, detail: str = "") -> None:
        FAIL.append(name)
        print(f"  FAIL  {name}" + (f"  ({detail})" if detail else ""))

    def section(title: str) -> None:
        print(f"\n=== {title} ===")

    def finish(suite: str) -> int:
        print(f"\n{suite}: {len(PASS)} passed, {len(FAIL)} failed")
        return 1 if FAIL else 0

    return ok, fail, section, finish


# ---- fixture ----

def copy_fixture() -> Path:
    """Fresh temp copy of workspace_mini — exercises mutate the COPY, never the
    checked-in fixture (so a full battery leaves the fixture byte-identical)."""
    dst = Path(tempfile.mkdtemp(prefix="ws_mini_"))
    shutil.copytree(FIXTURE, dst / "workspace_mini")
    ws = dst / "workspace_mini"
    for sub in ("_hq/meetings", "_hq/board-packs", "_hq/briefings"):
        (ws / sub).mkdir(parents=True, exist_ok=True)
    return ws


# ---- .docx extraction + normalization ----

def extract_docx_text(path: str | Path) -> str:
    """Extract text in document order: body paragraphs + table cells (row-joined
    with ' | '), then the page footer. Artifact paths are excluded (callers assert
    those separately) — this is structure + content, not file paths."""
    from docx import Document
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    doc = Document(str(path))
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            txt = Paragraph(child, doc).text.strip()
            if txt:
                lines.append(txt)
        elif isinstance(child, CT_Tbl):
            for row in Table(child, doc).rows:
                cells = [c.text.strip() for c in row.cells]
                lines.append(" | ".join(cells))
    # footer (page footer part, not body)
    try:
        for section in doc.sections:
            for p in section.footer.paragraphs:
                t = p.text.strip()
                if t:
                    lines.append(t)
    except Exception:
        pass
    return "\n".join(lines)


def normalize(text: str) -> str:
    """Dates -> <DATE>, collapse internal whitespace, drop blank lines, LF endings."""
    text = _DATE_ISO.sub("<DATE>", text)
    text = _DATE_LONG.sub("<DATE>", text)
    out = []
    for line in text.replace("\r\n", "\n").split("\n"):
        line = re.sub(r"\s+", " ", line).strip()
        if line:
            out.append(line)
    return "\n".join(out)


def assert_no_placeholders(text: str) -> list[str]:
    return [p for p in PLACEHOLDERS if p in text]


# ---- golden compare ----

def _update_mode(update: bool) -> bool:
    return update or os.environ.get("CR_UPDATE_GOLDENS") == "1"


def compare_golden(skill_name: str, text: str, update: bool = False) -> tuple[bool, str]:
    """Compare normalized `text` to tests/golden/<skill>.golden.txt. In update mode
    (arg or CR_UPDATE_GOLDENS=1) rewrites the golden and returns (True, '')."""
    GOLDEN_DIR.mkdir(parents=True, exist_ok=True)
    golden = GOLDEN_DIR / f"{skill_name}.golden.txt"
    norm = normalize(text).rstrip() + "\n"
    if _update_mode(update):
        golden.write_text(norm, encoding="utf-8")
        return True, ""
    if not golden.exists():
        return False, f"golden missing: {golden} — run CR_UPDATE_GOLDENS=1 to create it"
    want = golden.read_text(encoding="utf-8")
    if want == norm:
        return True, ""
    diff = "\n".join(difflib.unified_diff(
        want.splitlines(), norm.splitlines(),
        fromfile="golden", tofile="rendered", lineterm=""))
    return False, diff


# ---- event validator (reads the live schema enum) ----

_schema_cache = None


def _enum() -> set[str]:
    global _schema_cache
    if _schema_cache is None:
        _schema_cache = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    return set(_schema_cache["properties"]["type"]["enum"])


_THREAD_RE = re.compile(r"^project_[0-9]{3,}$")


def validate_event(ev: dict) -> list[str]:
    """Stdlib mini-validator over events.schema.json. Required keys present, `type`
    in the LIVE enum (so renaming a type the fixture uses fails immediately), seq is
    an int, primary_thread_id matches the pattern when present."""
    v: list[str] = []
    if not isinstance(ev, dict):
        return ["event is not an object"]
    for k in ("seq", "ts", "type", "source_skill", "data"):
        if k not in ev:
            v.append(f"missing required key '{k}'")
    t = ev.get("type")
    if t is not None and t not in _enum():
        v.append(f"type '{t}' not in the schema enum")
    if "seq" in ev and not isinstance(ev["seq"], int):
        v.append("seq is not an int")
    ptid = ev.get("primary_thread_id")
    if ptid is not None and not _THREAD_RE.match(str(ptid)):
        v.append(f"primary_thread_id '{ptid}' violates ^project_[0-9]{{3,}}$")
    return v
